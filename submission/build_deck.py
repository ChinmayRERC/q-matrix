#!/usr/bin/env python3
"""Build the submission deck from the organisers' light-purple template.

Run with a Python that has python-pptx installed:

    python3 submission/build_deck.py

The template and rendered circuit/results figures live next to this script.
The output is `submission/VTU-QFF-HC-0024.pptx`.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "template.pptx"
OUTPUT = ROOT / "VTU-QFF-HC-0024.pptx"
ASSETS = ROOT / "assets"

TEAM_CODE = "VTU-QFF-HC-0024"
DEEP = RGBColor(0x2E, 0x1B, 0x5B)
PURPLE = RGBColor(0x5B, 0x3B, 0xA8)
PINK = RGBColor(0xE8, 0x63, 0x9F)
MUTED = RGBColor(0x6F, 0x73, 0x99)
LIGHT_MUTED = RGBColor(0x8B, 0x86, 0xA8)
DARK = RGBColor(0x3A, 0x2D, 0x63)


def set_run(run, size=10, bold=False, color=DARK, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def clear_text(shape):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    return tf


def add_para(tf, first=False, space_before=0, space_after=2.5, line_spacing=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    return p


def header(tf, text, first=False, space_before=0, size=9.8):
    p = add_para(tf, first=first, space_before=space_before, space_after=2.0)
    r = p.add_run()
    r.text = text.upper()
    set_run(r, size=size, bold=True, color=PURPLE)
    return p


def bullet(tf, text, size=9.5, space_after=2.8):
    p = add_para(tf, space_after=space_after)
    b = p.add_run()
    b.text = "▸ "
    set_run(b, size=size, bold=True, color=PINK)
    r = p.add_run()
    r.text = text
    set_run(r, size=size, color=DARK)
    return p


def body_block(slide, shape_name, blocks, size=9.5, header_size=9.8, width=8.45, left=0.83, height=2.85):
    shape = next(s for s in slide.shapes if s.name == shape_name) if isinstance(shape_name, str) else shape_name
    shape.left = Inches(left)
    shape.top = Inches(2.03)
    shape.width = Inches(width)
    shape.height = Inches(height)
    tf = clear_text(shape)
    first = True
    for heading, lines in blocks:
        header(tf, heading, first=first, space_before=0 if first else 5.5, size=header_size)
        first = False
        for line in lines:
            bullet(tf, line, size=size)


def set_field(slide, name, value, size=11.5, hyperlink=None):
    shape = next(s for s in slide.shapes if s.name == name)
    tf = shape.text_frame
    # Field shapes have label (p0) and value (p1). Keep the label and replace
    # only the value paragraph with a fresh run.
    while len(tf.paragraphs) > 1:
        tf._element.remove(tf.paragraphs[-1]._p)
    if not tf.paragraphs or not tf.paragraphs[0].runs:
        p = tf.paragraphs[0] if tf.paragraphs else tf.add_paragraph()
    else:
        p = tf.add_paragraph()
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = 0.95
    r = p.add_run()
    r.text = value
    set_run(r, size=size, bold=False, color=DEEP)
    if hyperlink:
        r.hyperlink.address = hyperlink


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_picture(slide, filename):
    slide.shapes.add_picture(
        str(ASSETS / filename),
        Inches(0.5),
        Inches(1.85),
        width=Inches(5.4),
        height=Inches(3.2),
    )


def main():
    prs = Presentation(str(TEMPLATE))
    slides = list(prs.slides)

    prs.core_properties.title = "Q-matrix · Quantum Space"
    prs.core_properties.author = f"Team Q-matrix ({TEAM_CODE})"
    prs.core_properties.subject = "Qiskit Fall Fest 2026 · agent-reviewed quantum learning platform"
    prs.core_properties.keywords = "Qiskit, quantum education, BKT, NVIDIA, Q-matrix, Quantum Space"

    # ------------------------------------------------------------------
    # 1. Title / participant details
    # ------------------------------------------------------------------
    s = slides[0]
    set_field(s, "Field: PROJECT TITLE", "Q-matrix · Quantum Space", size=11.0)
    set_field(s, "Field: TRACK", "Quantum Education & Learning", size=9.3)
    set_field(s, "Field: TEAM NAME", "Q-matrix", size=10.5)
    set_field(s, "Field: TEAM LEADER", "Chinmay R", size=10.5)
    set_field(s, "Field: EMAIL / CONTACT", "github.com/ChinmayRERC", size=8.3, hyperlink="https://github.com/ChinmayRERC")
    set_field(s, "Field: TEAM MEMBERS", "Chinmay R  ·  M Naveen  ·  Thulasi Ram M", size=9.6)
    set_field(s, "Field: INSTITUTION / ORGANISATION", "Visvesvaraya Technological University", size=7.8)
    set_field(s, "Field: COUNTRY", "India", size=9.8)

    code = s.shapes.add_textbox(Inches(4.4), Inches(4.94), Inches(5.15), Inches(0.24))
    tf = code.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = f"TEAM CODE   {TEAM_CODE}"
    set_run(r, size=8.5, bold=True, color=PINK)

    add_notes(
        s,
        "Team code: VTU-QFF-HC-0024. Project: Q-matrix · Quantum Space. "
        "Public repository: https://github.com/vtu-qiskit-fall-fest-2026/vtu-hc-q-matrix. "
        "Live demo: https://chinmayrerc.github.io/q-matrix/.",
    )

    # ------------------------------------------------------------------
    # 2. Problem statement
    # ------------------------------------------------------------------
    s = slides[1]
    body_block(
        s,
        "Text 6",
        [
            (
                "The gap",
                [
                    "Quantum courses explain amplitudes, but learners cannot see them; notebooks and cloud queues add friction before the first insight.",
                    "Generic AI tutors hallucinate circuits and cannot read the learner's live state or mistake.",
                    "Instructors receive completion ticks, not evidence of conceptual change.",
                ],
            ),
            (
                "Who feels it",
                [
                    "Students in low-resource colleges without lab hardware or paid cloud credits.",
                    "Self-learners who stall between a video and a working Qiskit program.",
                    "Teachers who need a defensible mastery signal, not another quiz score.",
                ],
            ),
            (
                "Design goal",
                [
                    "One browser workspace: learn → build → simulate → export real Qiskit.",
                    "Grounded feedback from the circuit itself; every secret stays server-side.",
                ],
            ),
        ],
    )
    add_notes(s, "Problem framing from the 176-topic curriculum audit and the 7-stage lesson design. No learner PII is collected; progress is local-first.")

    # ------------------------------------------------------------------
    # 3. Objectives & novelty
    # ------------------------------------------------------------------
    s = slides[2]
    body_block(
        s,
        "Text 6",
        [
            (
                "Objectives",
                [
                    "Make quantum literacy measurable: a 7-stage loop from intuition to Qiskit export.",
                    "Ground every tutor answer in the live circuit, simulator state and static analysis.",
                    "Give educators an auditable mastery trail: BKT posterior, gain and challenge record.",
                ],
            ),
            (
                "Novelty",
                [
                    "Agent-reviewed content pipeline: deterministic analyzer + NVIDIA Nemotron tutor.",
                    "Challenge mode loads scaffolds, never solutions; Reset restores a clean start.",
                    "Browser-native density simulation with noise, mid-circuit measurement and feed-forward.",
                ],
            ),
            (
                "Scale",
                [
                    "176 topics · 20 domains · 8 learning paths · 13 circuit templates · 11 challenges.",
                    "147 automated checks, including 15 Qiskit 2.5.2 reference-algorithm tests.",
                ],
            ),
        ],
    )
    add_notes(s, "Metrics are counted from src/content/topics.generated.ts and src/lib/challenges.ts. Tests: 120 Vitest + 12 API + 15 Qiskit = 147.")

    # ------------------------------------------------------------------
    # 4. Solution & quantum advantage
    # ------------------------------------------------------------------
    s = slides[3]
    body_block(
        s,
        "Text 6",
        [
            (
                "What we built",
                [
                    "Q-matrix is the app, Quantum Space is the platform, Chinnu is the grounded tutor.",
                    "Learn → Compose → Practice → Galaxy → Progress: one continuous learning loop.",
                    "React SPA + FastAPI on Render; GitHub Pages demo; NVIDIA Nemotron 3.5 Lightning server-side.",
                ],
            ),
            (
                "Quantum advantage",
                [
                    "Real statevector/density simulation exposes superposition, interference, entanglement and measurement.",
                    "Noise and feed-forward are first-class: teleportation and mid-circuit correction run in-browser.",
                    "Every artifact is portable: OpenQASM 2.0, Qiskit, Cirq and PennyLane exports.",
                ],
            ),
            (
                "Why now",
                [
                    "Qiskit 2.5.2 maturity + browser compute + grounded LLMs = a real quantum classroom.",
                ],
            ),
        ],
    )
    add_notes(s, "The simulator is implemented in TypeScript; the backend mirrors the same statevector contract in NumPy for API parity. NVIDIA model: nvidia/nemotron-3.5-lightning-30b-a3b.")

    # ------------------------------------------------------------------
    # 5. Algorithm & circuit design
    # ------------------------------------------------------------------
    s = slides[4]
    for shape in list(s.shapes):
        if shape.name == "Text 5":
            shape._element.getparent().remove(shape._element)
    add_picture(s, "ghz_circuit.png")
    body_block(
        s,
        "Text 8",
        [
            (
                "Algorithm",
                [
                    "GHZ: H(q0), CNOT(q0,q1), CNOT(q1,q2) → (|000⟩+|111⟩)/√2.",
                    "Each column is one simultaneous moment; gates in a column commute.",
                ],
            ),
            (
                "Engine",
                [
                    "Statevector for ideal runs, density matrix for noise; measurement collapse + `if_test` feed-forward.",
                    "QASM parser reconstructs moments with per-qubit scheduling.",
                ],
            ),
            (
                "Evidence",
                [
                    "Bell/GHZ, Deutsch–Jozsa, Bernstein–Vazirani, Grover, QFT, QPE, teleportation, noisy Bell.",
                    "Web QFT matches QFTGate to 1e-12; Bell payload returns [0.5, 0, 0, 0.5].",
                ],
            ),
        ],
        size=8.35,
        header_size=8.7,
        width=2.95,
        left=6.33,
        height=2.95,
    )
    add_notes(s, "Circuit rendered with Qiskit 2.5.2. Reference implementations live in qiskit_ref/ and are exercised by 15 pytest cases.")

    # ------------------------------------------------------------------
    # 6. Implementation & tech stack
    # ------------------------------------------------------------------
    s = slides[5]
    body_block(
        s,
        "Text 6",
        [
            (
                "Frontend",
                [
                    "React 19 + Vite + TypeScript; custom SVG circuit board with keyboard placement and ARIA labels.",
                    "Hash routing keeps GitHub Pages deep links working without rewrites.",
                ],
            ),
            (
                "Quantum core",
                [
                    "TypeScript complex arithmetic; statevector and density simulation; QASM import/export.",
                    "13 templates; optimizer; static analyzer; Qiskit/Cirq/PennyLane exporters.",
                ],
            ),
            (
                "Backend & infra",
                [
                    "FastAPI proxy keeps NVIDIA_API_KEY server-side; deterministic fallback when offline.",
                    "Single-origin Docker service on Render; GitHub Actions deploys the static demo.",
                    "176-topic generated curriculum; BKT mastery and CSV/JSON educator export.",
                ],
            ),
        ],
    )
    add_notes(s, "Deployment: https://q-matrix.onrender.com (Docker) and https://chinmayrerc.github.io/q-matrix/ (static). The API key is never bundled in the client.")

    # ------------------------------------------------------------------
    # 7. Results & analysis
    # ------------------------------------------------------------------
    s = slides[6]
    for shape in list(s.shapes):
        if shape.name == "Text 5":
            shape._element.getparent().remove(shape._element)
    add_picture(s, "results.png")
    body_block(
        s,
        "Text 8",
        [
            (
                "What the data shows",
                [
                    "Ideal Bell: 4,095 / 4,097 counts for |00⟩ / |11⟩ in 8,192 shots.",
                    "Noise moves 7.7% of shots into |01⟩/|10⟩ and lowers purity; the UI surfaces entropy and concurrence.",
                ],
            ),
            (
                "Verification",
                [
                    "120 Vitest + 12 API + 15 Qiskit = 147/147 passing.",
                    "QFT vs QFTGate: max error 1.2e-12; TypeScript strict; production build clean.",
                ],
            ),
            (
                "Live",
                [
                    "API health: numpy-statevector + NVIDIA; tutor verified end-to-end on Pages.",
                    "Bell API returns [0.5, 0, 0, 0.5].",
                ],
            ),
        ],
        size=8.35,
        header_size=8.7,
        width=2.95,
        left=6.33,
        height=2.95,
    )
    add_notes(s, "The left panel is an actual Aer run with depolarising + readout error (seed 7). The right panel reports the repository's automated checks, not a projected metric.")

    # ------------------------------------------------------------------
    # 8. Impact & applications
    # ------------------------------------------------------------------
    s = slides[7]
    body_block(
        s,
        "Text 6",
        [
            (
                "For learners",
                [
                    "Zero install, low bandwidth; first simulated circuit in under a minute.",
                    "Immediate, grounded feedback when a gate is misplaced or a state is separable.",
                ],
            ),
            (
                "For educators",
                [
                    "Class-level mastery map, pretest/posttest, normalised gain and challenge record.",
                    "CSV/JSON export for spreadsheets or LMS import; no database required.",
                ],
            ),
            (
                "For institutions",
                [
                    "Low-cost quantum lab on existing laptops; no GPU or cloud queue.",
                    "Reproducible Qiskit artifacts bridge classroom work to real SDK workflows.",
                ],
            ),
            (
                "For industry",
                [
                    "Faster onboarding for quantum software teams; a reusable simulator/analysis core.",
                ],
            ),
        ],
    )
    add_notes(s, "Impact hypothesis: a browser-native, grounded lab removes the two biggest drop-off points we observed: environment setup and ungrounded feedback.")

    # ------------------------------------------------------------------
    # 9. Challenges & future scope
    # ------------------------------------------------------------------
    s = slides[8]
    body_block(
        s,
        "Text 6",
        [
            (
                "What is hard",
                [
                    "Browser compute ceilings; density simulation is expensive beyond small registers.",
                    "Approximate noise models; LLM latency and provider availability.",
                    "Content authoring quality; multi-user persistence is not yet included.",
                ],
            ),
            (
                "How we mitigate",
                [
                    "Statevector cap + density path; deterministic fallback; server-side key; localStorage.",
                    "Content generator + agent review; explicit limitations in the README.",
                ],
            ),
            (
                "Next",
                [
                    "Qiskit Runtime/hardware labs, adaptive BKT + spaced repetition, LMS/LTI integration.",
                    "WebGPU simulator, VQE/QAOA/Shor modules, multilingual Chinnu, classroom pilots.",
                ],
            ),
        ],
    )
    add_notes(s, "The roadmap prioritises hardware-backed labs and classroom pilots; the current build is intentionally local-first and single-user.")

    # ------------------------------------------------------------------
    # 10. Closing
    # ------------------------------------------------------------------
    s = slides[9]
    set_field(
        s,
        "Field: GITHUB REPOSITORY",
        "vtu-qiskit-fall-fest-2026/vtu-hc-q-matrix",
        size=6.6,
        hyperlink="https://github.com/vtu-qiskit-fall-fest-2026/vtu-hc-q-matrix",
    )
    # Rename the demo field so the live link is unambiguous.
    demo_label = next(shape for shape in s.shapes if shape.name == "Field: DEMO VIDEO LINK")
    demo_label.text_frame.paragraphs[0].runs[0].text = "LIVE DEMO / VIDEO"
    set_field(
        s,
        "Field: DEMO VIDEO LINK",
        "chinmayrerc.github.io/q-matrix",
        size=8.3,
        hyperlink="https://chinmayrerc.github.io/q-matrix/",
    )

    links = s.shapes.add_textbox(Inches(0.5), Inches(4.34), Inches(5.4), Inches(0.78))
    tf = links.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_top = 0
    p = tf.paragraphs[0]
    p.space_after = Pt(1.5)
    r = p.add_run()
    r.text = "Live demo · https://chinmayrerc.github.io/q-matrix"
    set_run(r, size=8.0, color=DEEP)
    r.hyperlink.address = "https://chinmayrerc.github.io/q-matrix/"
    p2 = tf.add_paragraph()
    p2.space_after = Pt(1.5)
    r2 = p2.add_run()
    r2.text = "API · https://q-matrix.onrender.com"
    set_run(r2, size=8.0, color=DEEP)
    r2.hyperlink.address = "https://q-matrix.onrender.com"
    p3 = tf.add_paragraph()
    r3 = p3.add_run()
    r3.text = f"Team code · {TEAM_CODE}"
    set_run(r3, size=8.0, bold=True, color=PINK)
    add_notes(s, "Public repo: https://github.com/vtu-qiskit-fall-fest-2026/vtu-hc-q-matrix. Live demo: https://chinmayrerc.github.io/q-matrix/. API: https://q-matrix.onrender.com.")

    prs.save(str(OUTPUT))
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
